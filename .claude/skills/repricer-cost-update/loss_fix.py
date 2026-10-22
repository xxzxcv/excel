"""Loss-prevention fix for a repricer export (Banner SKUs only; in-house rows are never touched).

Cost
  Sim price file       cost = round(New Price ex VAT x pack qty, 2)
  Banner catalogue     cost = max(current cost, round(Price Web x 1.20 x pack qty, 2))   (flat 20% increase)
  neither              current cost kept
Shipping (Banner charges 3.50 per order line under 16.50; existing shipping is never removed)
  Sim price            3.50 if cost < 16.50 else 0
  catalogue only       3.50 if Price Web x qty < 16.50 (catches both the old and the +20% price); 0 only if
                       even the old price clears 16.50
  neither              only ever added (3.50 if current cost < 16.50), never removed
VAT
  0 -> 20 unless Amazon's tax category (from Sellerboard) says the product is not standard rated.
Raw lines are copied; only unit_cost, shipping_cost and vat_rate are replaced. Rows that need no change are
left out of the upload files.
"""
import csv, os, re, sys, collections
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import load_workbook
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import repricer_cost_update as R

src, sim_path, sb_path, catalogue_path, outdir = sys.argv[1:6]
taxcat_paths = sys.argv[6:]
T, FEE, ZERO, UPLIFT = Decimal('16.50'), Decimal('3.50'), Decimal('0'), Decimal('1.20')
# Sellerboard shows these at the Banner price, so the "Banner price far below current cost" hold does not apply
CONFIRMED_LOW = {'0717366'}
# Amazon UK product tax categories that are not the 20% standard rate
NON_STANDARD = {'A_BOOKS_GEN': '0', 'A_FOOD_GEN': '0', 'A_GEN_NOTAX': '0', 'A_GEN_ZERO': '0',
                'A_GEN_REDUCED': '5', 'A_FOOD_TEA': '0', 'A_FOOD_PLAINBISCUIT': '0', 'A_FOOD_CEREALBARS': '0',
                'A_FOOD_SWEETNR': '0', 'A_CLTH_CHILD': '0', 'A_CLTH_BABY': '0'}

# Amazon files these under A_FOOD_TEA but they are (or may be) standard-rated in the UK: nappies, chocolate
# biscuits, plant-based drinks, ready-to-drink coffee. 20% is the safe choice.
FORCE_STANDARD = {'8009083', '8914627', '7900618', '7900672', '7469086', '7459134', '7508490', '7487638',
                  '3297929', '3292947'}
sim = R.load_sim(sim_path); codes = set(sim); labels = R.load_labels(sb_path)
ws = load_workbook(catalogue_path, read_only=True, data_only=True)['Catalogue']
catalogue = {}
for row in ws.iter_rows(min_row=2, max_col=14, values_only=True):
    if row[0] is not None and row[10] is not None:
        code = str(row[0]).strip().upper()
        # column N holds the pre-21/10 price for codes the Sim file updated; K is the price otherwise
        catalogue[code] = Decimal(str(row[10]))
cat_codes = set(catalogue) - codes
taxcat = {}
for p in taxcat_paths:
    for r in csv.DictReader(open(p, encoding='utf-8')):
        taxcat.setdefault(r['sku'].strip(), set()).add(r['vat_category'])

TOKEN = re.compile(r'"(?:[^"]|"")*"|[^,]*')
def split_raw(line):
    toks, pos = [], 0
    while True:
        m = TOKEN.match(line, pos); toks.append(m.group(0)); pos = m.end()
        if pos >= len(line): return toks
        assert line[pos] == ','; pos += 1

def money(x):
    return x.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

lines = open(src, encoding='utf-8', newline='').read().split('\n'); assert lines[-1] == ''
hdr = next(csv.reader([lines[0]]))
UC, SC, VC = hdr.index('unit_cost'), hdr.index('shipping_cost'), hdr.index('vat_rate')
zero_up, rest_up, audit, stats = [lines[0]], [lines[0]], [], collections.Counter()
for line in lines[1:-1]:
    f = next(csv.reader([line])); toks = split_raw(line)
    assert len(toks) == len(hdr) == len(f) and next(csv.reader([','.join(toks)])) == f
    sku = f[0]; label = labels.get(sku.strip(), 'NOT IN SELLERBOARD')
    cur_c, cur_s, cur_v = Decimal(f[UC] or 0), Decimal(f[SC] or 0), f[VC]
    rec = dict(sku=sku, merchant_id=f[2], asin=f[5], title=f[4][:90], label=label, source='', supplier_code='',
               pack_qty='', old_cost=f[UC], new_cost='', old_shipping=f[SC], new_shipping='', old_vat=cur_v,
               new_vat='', changed='', note='')
    group = 'zero' if sku.startswith('0') else 'rest'

    def skip(note, key):
        rec['note'] = note; rec['changed'] = 'not touched'; audit.append(rec); stats[f'{group}: {key}'] += 1

    if label == 'INHOUSE':
        skip('In-house', 'not touched - in-house'); continue
    code, mult, how = R.parse_sku(sku, codes)
    source = 'sim' if code else ''
    if not code:
        code, mult, how = R.parse_sku(sku, cat_codes)
        source = 'catalogue' if code else ''
    if not code:
        loose = R.loose_code(sku, codes) or R.loose_code(sku, cat_codes)
        if loose:
            code, mult = loose, 1; source = 'sim' if loose in codes else 'catalogue'
            rec['note'] = 'Code inside SKU text, pack qty 1 (approved 21/10)'
    if not code and label != 'BANNER':
        skip('Not a Banner SKU', 'not touched - not Banner'); continue
    rec.update(source=source or 'current cost', supplier_code=code or '', pack_qty=mult or '')

    if source == 'sim':
        old_line = sim[code][1] * mult
        if cur_c > 0 and cur_c / old_line > Decimal('2.5') and sku not in CONFIRMED_LOW:
            skip(f'Banner price £{money(sim[code][2] * mult)} is far below current cost £{cur_c}: check', 'held - price check')
            continue
        cost = money(sim[code][2] * mult)
        ship = ZERO if cost >= T else FEE
    elif source == 'catalogue':
        pw = catalogue[code] * mult
        floor = money(pw * UPLIFT)
        cost = max(cur_c, floor)
        ship = ZERO if pw >= T else FEE
        if cost != cur_c:
            rec['note'] = f'Not in Sim file: Banner Price Web £{money(pw)} + 20% = £{floor}'
    else:
        if cur_c <= 0:
            skip('Banner, no cost and no supplier price', 'not touched - no cost'); continue
        cost = cur_c
        ship = max(cur_s, ZERO if cost >= T else FEE)

    # Loss prevention: never take away shipping that is already there (removing it would lower the floor price)
    ship = max(ship, cur_s)

    cats = taxcat.get(sku.strip(), set())
    if sku.strip() in FORCE_STANDARD:
        cats = {'A_GEN_STANDARD'}
    if cur_v in ('0', ''):
        non_std = {NON_STANDARD.get(c) for c in cats if c != 'A_GEN_STANDARD'}
        if not cats or 'A_GEN_STANDARD' in cats and not non_std:
            vat = '20'
        elif len(cats) == 1 and None not in non_std:
            vat = non_std.pop()
        else:
            vat = '20'
            rec['note'] = (rec['note'] + '; ' if rec['note'] else '') + f'Unusual tax category {sorted(cats)}: set to 20 to be safe'
    else:
        vat = cur_v

    changes = []
    if cost != cur_c: changes.append('cost')
    if ship != cur_s: changes.append('shipping added' if ship > cur_s else 'shipping removed')
    if vat != cur_v: changes.append(f'VAT {cur_v or "blank"}->{vat}')
    rec.update(new_cost=f'{cost:.2f}', new_shipping=f'{ship:.2f}', new_vat=vat,
               changed=' + '.join(changes) or 'already correct')
    audit.append(rec)
    for c in changes or ['already correct']:
        stats[f'{group}: {c.split(" ")[0] if c.startswith("VAT") else c}'] += 1
    if changes:
        toks[UC], toks[SC], toks[VC] = f'{cost:.2f}', f'{ship:.2f}', vat
        (zero_up if group == 'zero' else rest_up).append(','.join(toks))
        stats[f'{group}: rows in upload'] += 1

for name, rows in (('Repricer_Upload_1_SKUs_Starting_With_0.csv', zero_up),
                   ('Repricer_Upload_2_All_Other_Banner_SKUs.csv', rest_up)):
    open(f'{outdir}/{name}', 'w', encoding='utf-8', newline='').write('\n'.join(rows) + '\n')
with open(f'{outdir}/Fix_Audit_All_SKUs.csv', 'w', encoding='utf-8-sig', newline='') as fh:
    w = csv.DictWriter(fh, fieldnames=list(audit[0])); w.writeheader(); w.writerows(audit)
for k, v in sorted(stats.items()):
    print(f'{k:<40} {v}')
