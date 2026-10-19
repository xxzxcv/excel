#!/usr/bin/env python3
"""Calculate the daily Banner dropship payment from an orders CSV.

Same rules as Banner_Payment_Master_Template.xlsx:
  unit price = catalogue 'Price Web ' (col K) for the line's Product Code (col A)
  base       = round(unit price * quantity, 2)
  postage    = 0 if base >= 16.50 else 3.50   (per line)
  line total = base + postage;  grand total = sum of line totals
Codes not found exactly are retried with 1-3 leading zeros (the order export strips them).

Usage: banner_payment.py ORDERS.csv [--catalogue FILE.xlsx] [--lines OUT.csv]
"""
import argparse, csv, os, sys
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_CATALOGUE = os.path.normpath(os.path.join(HERE, "..", "..", "..", "Banner_Payment_Master_Template.xlsx"))
THRESHOLD, FEE = 16.50, 3.50


def load_catalogue(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    # Banner's own file uses 'Full File ', the master template uses 'Catalogue'
    names = {n.strip().lower(): n for n in wb.sheetnames}
    ws = wb[names.get("catalogue") or names["full file"]]
    prices = {}
    for row in ws.iter_rows(min_row=2, max_col=11, values_only=True):
        if row[0] is not None and row[10] is not None:
            prices[str(row[0]).strip()] = (float(row[10]), row[1])
    return prices


def lookup(prices, code):
    for pad in ("", "0", "00", "000"):
        if pad + code in prices:
            return pad + code, prices[pad + code]
    return None, (None, None)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("orders")
    ap.add_argument("--catalogue", default=DEFAULT_CATALOGUE)
    ap.add_argument("--lines", help="write a per-line breakdown CSV here")
    a = ap.parse_args()

    prices = load_catalogue(a.catalogue)
    with open(a.orders, newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.reader(fh))
    header, data = rows[0], [r for r in rows[1:] if len(r) > 11 and r[10].strip()]

    out, problems = [], []
    goods = postage = 0.0
    items = post_lines = 0
    for r in data:
        code, qty_s = r[10].strip(), r[11].strip()
        matched, (price, desc) = lookup(prices, code)
        try:
            qty = float(qty_s)
        except ValueError:
            qty = None
        if matched is None or qty is None or qty <= 0:
            status = "CODE NOT IN CATALOGUE" if matched is None else "CHECK QUANTITY"
            problems.append((r[0], code, qty_s, status))
            out.append([r[0], code, matched or "", qty_s, desc or "", "", "", "", "", status])
            continue
        base = round(price * qty, 2)
        post = 0.0 if base >= THRESHOLD else FEE
        goods += base; postage += post; items += qty; post_lines += post > 0
        out.append([r[0], code, matched, qty_s, desc, f"{price:.2f}", f"{base:.2f}",
                    f"{post:.2f}", f"{base + post:.2f}", "OK"])

    print(f"Orders file:            {os.path.basename(a.orders)}")
    print(f"Order lines:            {len(data)}  (priced {len(data) - len(problems)}, NOT priced {len(problems)})")
    print(f"Total items:            {items:g}")
    print(f"Goods value:            £{goods:,.2f}")
    print(f"Postage:                £{postage:,.2f}  ({post_lines} lines under £{THRESHOLD:.2f})")
    print(f"GRAND TOTAL TO PAY:     £{goods + postage:,.2f}")
    padded = [(o[1], o[2]) for o in out if o[2] and o[1] != o[2]]
    if padded:
        print(f"Leading zeros restored: {len(padded)} line(s): " + ", ".join(sorted({f'{c}->{m}' for c, m in padded})))
    if problems:
        print("WARNING - lines NOT included in the total:")
        for p in problems:
            print(f"  order {p[0]}  code {p[1]}  qty {p[2]}  -> {p[3]}")

    if a.lines:
        with open(a.lines, "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            w.writerow(["Order Number", "Product Code (file)", "Catalogue Code", "Quantity", "Description",
                        "Unit Price", "Base Line Value", "Postage", "Final Line Total", "Status"])
            w.writerows(out)
            w.writerow([]); w.writerow(["", "", "", "", "GRAND TOTAL", "", f"{goods:.2f}", f"{postage:.2f}", f"{goods + postage:.2f}"])
        print(f"Line breakdown written: {a.lines}")


if __name__ == "__main__":
    sys.exit(main())
